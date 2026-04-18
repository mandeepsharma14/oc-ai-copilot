"""
Document ingestion pipeline.
Parse → chunk → embed → upsert into Azure AI Search index.
"""
from __future__ import annotations
import uuid
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".docx", ".xlsx", ".txt", ".md"}


@dataclass
class DocumentChunk:
    chunk_id:      str
    content:       str
    content_vector: list
    document_id:   str
    document_name: str
    section:       Optional[str]
    version:       Optional[str]
    approved_date: Optional[str]
    domain:        Optional[str]
    product_line:  Optional[str]
    site:          Optional[str]
    stream_type:   str


class IngestionPipeline:
    """End-to-end document ingestion: parse → chunk → embed → index."""

    def __init__(self):
        self._embed_client = None

    async def ingest(
        self,
        file_path:   str,
        stream_type: str,
        metadata:    dict,
        index_name:  str,
    ) -> dict:
        path = Path(file_path)
        if path.suffix.lower() not in SUPPORTED:
            raise ValueError(f"Unsupported file type: {path.suffix}")

        doc_id = str(uuid.uuid4())
        logger.info(f"Ingesting {path.name} → {index_name}")

        text_blocks = self._parse(path)
        chunks_text = self._chunk(text_blocks)
        doc_chunks  = []

        for i, text in enumerate(chunks_text):
            vector = await self._embed(text)
            doc_chunks.append(DocumentChunk(
                chunk_id=f"{doc_id}_{i}",
                content=text,
                content_vector=vector,
                document_id=doc_id,
                document_name=metadata.get("name", path.name),
                section=metadata.get("section"),
                version=metadata.get("version"),
                approved_date=metadata.get("approved_date"),
                domain=metadata.get("domain"),
                product_line=metadata.get("product_line"),
                site=metadata.get("site"),
                stream_type=stream_type,
            ))

        await self._upsert(doc_chunks, index_name)

        return {
            "document_id":   doc_id,
            "document_name": metadata.get("name", path.name),
            "chunks_created": len(doc_chunks),
            "stream_type":   stream_type,
            "status":        "indexed",
        }

    def _parse(self, path: Path) -> list[str]:
        suffix = path.suffix.lower()
        try:
            if suffix in {".txt", ".md"}:
                return [path.read_text(encoding="utf-8")]
            elif suffix == ".docx":
                import docx
                doc = docx.Document(str(path))
                return [p.text for p in doc.paragraphs if p.text.strip()]
            elif suffix == ".pdf":
                import fitz
                doc = fitz.open(str(path))
                return [page.get_text() for page in doc if page.get_text().strip()]
            elif suffix == ".xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
                blocks = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        line = " | ".join(str(c) for c in row if c is not None)
                        if line.strip():
                            blocks.append(line)
                return blocks
        except Exception as e:
            logger.error(f"Parse error for {path.name}: {e}")
        return []

    def _chunk(
        self,
        blocks:     list[str],
        chunk_size: int = 512,
        overlap:    int = 102,
    ) -> list[str]:
        from app.core.config import settings
        chunk_size = settings.RAG_CHUNK_SIZE
        overlap    = settings.RAG_CHUNK_OVERLAP

        full_text = " ".join(b for b in blocks if b.strip())
        words     = full_text.split()
        chunks, i = [], 0
        while i < len(words):
            chunk = " ".join(words[i : i + chunk_size])
            if len(chunk.strip()) > 50:
                chunks.append(chunk)
            i += chunk_size - overlap
        return chunks

    async def _embed(self, text: str) -> list:
        from app.core.config import settings
        if not settings.AZURE_OPENAI_ENDPOINT:
            return []
        try:
            if self._embed_client is None:
                from openai import AsyncAzureOpenAI
                self._embed_client = AsyncAzureOpenAI(
                    azure_endpoint=settings.AZURE_OPENAI_ENDPOINT,
                    api_key=settings.AZURE_OPENAI_API_KEY,
                    api_version=settings.AZURE_OPENAI_API_VERSION,
                )
            resp = await self._embed_client.embeddings.create(
                input=text[:8000],
                model=settings.AZURE_OPENAI_DEPLOYMENT_EMBED,
            )
            return resp.data[0].embedding
        except Exception as e:
            logger.error(f"Embedding error: {e}")
            return []

    async def _upsert(self, chunks: list[DocumentChunk], index_name: str) -> None:
        from app.core.config import settings
        if not settings.AZURE_SEARCH_ENDPOINT:
            logger.info(f"Azure Search not configured — skipping upsert of {len(chunks)} chunks")
            return
        try:
            from azure.search.documents.aio import SearchClient
            from azure.core.credentials import AzureKeyCredential
            docs = [
                {
                    "chunk_id":       c.chunk_id,
                    "content":        c.content,
                    "content_vector": c.content_vector,
                    "document_id":    c.document_id,
                    "document_name":  c.document_name,
                    "section":        c.section,
                    "version":        c.version,
                    "approved_date":  c.approved_date,
                    "domain":         c.domain,
                    "product_line":   c.product_line,
                    "site":           c.site,
                    "stream_type":    c.stream_type,
                }
                for c in chunks
            ]
            async with SearchClient(
                endpoint=settings.AZURE_SEARCH_ENDPOINT,
                index_name=index_name,
                credential=AzureKeyCredential(settings.AZURE_SEARCH_ADMIN_KEY),
            ) as client:
                await client.upload_documents(documents=docs)
            logger.info(f"Upserted {len(docs)} chunks to {index_name}")
        except Exception as e:
            logger.error(f"Upsert error: {e}")
